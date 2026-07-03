import os
import re
import glob
import camelot
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

FLAVORS = ["lattice", "stream", "network", "hybrid", "ml", "auto"]
BACKENDS = ["pdfium", "ghostscript", "poppler"]
TABLE_STYLE_PRESETS = {
    "Blue / White": "TableStyleMedium2",
    "Light Blue": "TableStyleLight9",
    "Grey / White": "TableStyleMedium6",
    "Green / White": "TableStyleMedium7",
    "None (plain)": None,
}


def collect_pdf_paths(input_path, recursive=False):
    if os.path.isfile(input_path):
        return [input_path]
    pattern = "**/*.pdf" if recursive else "*.pdf"
    paths = glob.glob(os.path.join(input_path, pattern), recursive=recursive)
    return sorted(paths)


def get_page_count(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        return len(pdf.pages)


def pages_with_text(pdf_path, search_text, case_sensitive=False):
    matches = []
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            haystack = text if case_sensitive else text.lower()
            needle = search_text if case_sensitive else search_text.lower()
            if needle in haystack:
                matches.append(i)
    return matches


def group_words_into_lines(words, tol=3):
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        placed = False
        for line in lines:
            if abs(line["top"] - w["top"]) <= tol:
                line["words"].append(w)
                line["top"] = min(line["top"], w["top"])
                line["bottom"] = max(line["bottom"], w["bottom"])
                line["x0"] = min(line["x0"], w["x0"])
                line["x1"] = max(line["x1"], w["x1"])
                placed = True
                break
        if not placed:
            lines.append({
                "top": w["top"], "bottom": w["bottom"],
                "x0": w["x0"], "x1": w["x1"], "words": [w],
            })
    for line in lines:
        line["words"].sort(key=lambda w: w["x0"])
        line["text"] = " ".join(w["text"] for w in line["words"])
    return lines


def find_text_line_bbox(page, search_text, case_sensitive=False):
    words = page.extract_words()
    lines = group_words_into_lines(words)
    needle = search_text if case_sensitive else search_text.lower()
    best = None
    for line in lines:
        haystack = line["text"] if case_sensitive else line["text"].lower()
        if needle in haystack:
            best = line
            break
    return best


def derive_table_area(page, header_text, footer_text, case_sensitive=False, padding=4):
    page_height = float(page.height)
    page_width = float(page.width)

    top_pdfplumber = 0.0
    bottom_pdfplumber = page_height

    if header_text:
        line = find_text_line_bbox(page, header_text, case_sensitive)
        if line is not None:
            top_pdfplumber = line["bottom"] + padding

    if footer_text:
        line = find_text_line_bbox(page, footer_text, case_sensitive)
        if line is not None:
            bottom_pdfplumber = line["top"] - padding

    if bottom_pdfplumber <= top_pdfplumber:
        return None

    x1 = 0.0
    y1 = page_height - top_pdfplumber
    x2 = page_width
    y2 = page_height - bottom_pdfplumber
    return f"{x1:.1f},{y1:.1f},{x2:.1f},{y2:.1f}"


def build_per_page_areas(pdf_path, page_numbers, header_text, footer_text, case_sensitive=False):
    per_page = {}
    with pdfplumber.open(pdf_path) as pdf:
        for pno in page_numbers:
            if pno - 1 >= len(pdf.pages):
                continue
            page = pdf.pages[pno - 1]
            area = derive_table_area(page, header_text, footer_text, case_sensitive)
            if area:
                per_page[pno] = {"table_areas": [area]}
    return per_page


def parse_strip_text(raw):
    if not raw:
        return ""
    return raw.replace("\\n", "\n").replace("\\t", "\t")


def parse_replace_text(raw):
    mapping = {}
    if not raw:
        return None
    for line in raw.splitlines():
        line = line.strip()
        if not line or "=>" not in line:
            continue
        old, new = line.split("=>", 1)
        mapping[old] = new
    return mapping or None


def parse_area_list(raw):
    if not raw:
        return None
    areas = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    return areas or None


def parse_columns_list(raw):
    if not raw:
        return None
    cols = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    return cols or None


def build_read_kwargs(cfg, flavor):
    kwargs = {}

    if flavor in ("lattice", "hybrid", "ml"):
        kwargs["backend"] = cfg["backend"]
        kwargs["use_fallback"] = cfg["use_fallback"]

    if flavor in ("lattice", "hybrid"):
        kwargs["engine"] = cfg["engine"]

    if cfg["strip_text"]:
        kwargs["strip_text"] = parse_strip_text(cfg["strip_text"])

    replace_map = parse_replace_text(cfg["replace_text"])
    if replace_map:
        kwargs["replace_text"] = replace_map

    if cfg["split_text"]:
        kwargs["split_text"] = True
    if cfg["flag_size"]:
        kwargs["flag_size"] = True

    if flavor != "ml":
        areas = parse_area_list(cfg["table_areas"])
        if areas:
            kwargs["table_areas"] = areas
        regions = parse_area_list(cfg["table_regions"])
        if regions:
            kwargs["table_regions"] = regions
        cols = parse_columns_list(cfg["columns"])
        if cols and flavor in ("stream", "network", "hybrid"):
            kwargs["columns"] = cols

    if flavor in ("stream", "network", "hybrid"):
        kwargs["row_tol"] = cfg["row_tol"]
        kwargs["column_tol"] = cfg["column_tol"]
        if cfg["edge_tol"] is not None:
            kwargs["edge_tol"] = cfg["edge_tol"]

    if flavor in ("lattice", "hybrid"):
        kwargs["line_scale"] = cfg["line_scale"]
        kwargs["process_background"] = cfg["process_background"]
        kwargs["line_tol"] = cfg["line_tol"]
        kwargs["joint_tol"] = cfg["joint_tol"]
        kwargs["threshold_blocksize"] = cfg["threshold_blocksize"]
        kwargs["threshold_constant"] = cfg["threshold_constant"]
        kwargs["iterations"] = cfg["iterations"]
        kwargs["erode_iterations"] = cfg["erode_iterations"]
        kwargs["resolution"] = cfg["resolution"]

    if flavor == "ml":
        kwargs["structure_model"] = cfg["structure_model"]
        kwargs["detection_model"] = cfg["detection_model"]
        kwargs["device"] = cfg["device"]
        kwargs["detection_threshold"] = cfg["detection_threshold"]
        kwargs["structure_threshold"] = cfg["structure_threshold"]
        kwargs["crop_padding"] = cfg["crop_padding"]
        kwargs["ocr"] = cfg["ocr"]
        kwargs["resolution"] = cfg["resolution"]
        kwargs["backend"] = cfg["backend"]
        kwargs["use_fallback"] = cfg["use_fallback"]

    if cfg["password"]:
        kwargs["password"] = cfg["password"]

    if cfg["parallel"]:
        kwargs["parallel"] = True
        if cfg["cpu_count"]:
            kwargs["cpu_count"] = cfg["cpu_count"]

    return kwargs


def extract_tables_for_pdf(pdf_path, cfg, log):
    total_pages = get_page_count(pdf_path)

    if cfg["filter_pages_enabled"] and cfg["filter_text"]:
        matched = pages_with_text(pdf_path, cfg["filter_text"], cfg["filter_case_sensitive"])
        if not matched:
            log(f"  no pages matched filter text in {os.path.basename(pdf_path)}, skipping")
            return []
        pages_str = ",".join(str(p) for p in matched)
        log(f"  pages kept after text filter: {pages_str}")
    else:
        pages_str = cfg["pages"] or "all"
        if pages_str.lower() == "all":
            matched = list(range(1, total_pages + 1))
        else:
            matched = None

    per_page = None
    if cfg["derive_region_enabled"] and (cfg["header_text"] or cfg["footer_text"]):
        target_pages = matched if matched else list(range(1, total_pages + 1))
        per_page = build_per_page_areas(
            pdf_path, target_pages, cfg["header_text"], cfg["footer_text"],
            cfg["derive_case_sensitive"],
        )
        if per_page:
            log(f"  derived table regions for {len(per_page)} page(s) from header/footer text")

    flavors_to_run = FLAVORS if cfg["try_all_flavors"] else [cfg["flavor"]]

    results = []
    for flavor in flavors_to_run:
        kwargs = build_read_kwargs(cfg, flavor)
        if per_page:
            kwargs["per_page"] = per_page
        try:
            tables = camelot.read_pdf(pdf_path, pages=pages_str, flavor=flavor, **kwargs)
        except Exception as e:
            log(f"  flavor='{flavor}' failed: {e}")
            continue

        if tables.n == 0:
            log(f"  flavor='{flavor}': 0 tables found")
            continue

        if cfg["stitch_enabled"]:
            before = tables.n
            tables = tables.stack_contiguous(
                match=cfg["stitch_match"], keep_first_header=cfg["stitch_keep_first_header"],
            )
            log(f"  flavor='{flavor}': {before} physical -> {tables.n} logical table(s) after stitching")
        else:
            log(f"  flavor='{flavor}': {tables.n} table(s)")

        for t in tables:
            results.append({
                "file": os.path.basename(pdf_path),
                "flavor": flavor,
                "backend": kwargs.get("backend", ""),
                "page": t.page,
                "shape": t.shape,
                "accuracy": t.accuracy,
                "whitespace": t.whitespace,
                "df": t.df,
            })

    return results


def sanitize_sheet_name(name, existing):
    name = re.sub(r"[\[\]:*?/\\]", "_", name)[:31]
    base = name
    i = 1
    while name in existing:
        suffix = f"_{i}"
        name = base[: 31 - len(suffix)] + suffix
        i += 1
    existing.add(name)
    return name


def autofit_columns(ws, start_row, start_col, n_rows, n_cols, min_width=8, max_width=60):
    for c in range(n_cols):
        col_letter = get_column_letter(start_col + c)
        longest = min_width
        for r in range(n_rows):
            val = ws.cell(row=start_row + r, column=start_col + c).value
            if val is not None:
                longest = max(longest, len(str(val)))
        ws.column_dimensions[col_letter].width = min(longest + 2, max_width)


def write_table_sheet(wb, sheet_name, record, style_cfg, existing_names):
    name = sanitize_sheet_name(sheet_name, existing_names)
    ws = wb.create_sheet(title=name)

    base_font = Font(name=style_cfg["font_name"], size=style_cfg["font_size"])
    meta_font = Font(name=style_cfg["font_name"], size=style_cfg["font_size"], bold=True)

    meta_rows = [
        ("File", record["file"]),
        ("Page", record["page"]),
        ("Flavor", record["flavor"]),
        ("Backend", record["backend"] or "n/a"),
        ("Accuracy", f"{record['accuracy']:.1f}"),
        ("Whitespace", f"{record['whitespace']:.1f}"),
    ]
    for i, (k, v) in enumerate(meta_rows, start=1):
        ws.cell(row=i, column=1, value=k).font = meta_font
        ws.cell(row=i, column=2, value=v).font = base_font

    header_row = len(meta_rows) + 2
    df = record["df"]
    n_rows, n_cols = df.shape

    for c in range(n_cols):
        col_name = str(df.iloc[0, c]) if n_rows > 0 else f"col_{c + 1}"
        cell = ws.cell(row=header_row, column=c + 1, value=col_name)
        cell.font = Font(name=style_cfg["font_name"], size=style_cfg["font_size"], bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=style_cfg["header_color"], end_color=style_cfg["header_color"], fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in range(1, n_rows):
        for c in range(n_cols):
            val = df.iloc[r, c]
            cell = ws.cell(row=header_row + r, column=c + 1, value=val)
            cell.font = base_font

    data_last_row = header_row + max(n_rows - 1, 0)
    if data_last_row >= header_row and n_cols > 0:
        ref = f"{get_column_letter(1)}{header_row}:{get_column_letter(n_cols)}{data_last_row}"
        table_style_name = TABLE_STYLE_PRESETS.get(style_cfg["style_preset"])
        if table_style_name:
            tbl = Table(displayName=f"Tbl_{name}".replace(" ", "_")[:30], ref=ref)
            tbl.tableStyleInfo = TableStyleInfo(
                name=table_style_name, showRowStripes=True, showColumnStripes=False,
                showFirstColumn=False, showLastColumn=False,
            )
            ws.add_table(tbl)

    if style_cfg["autofit"]:
        autofit_columns(ws, header_row, 1, n_rows, n_cols)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    return name


def write_summary_sheet(wb, records, sheet_names, style_cfg):
    ws = wb.create_sheet(title="Summary", index=0)
    base_font = Font(name=style_cfg["font_name"], size=style_cfg["font_size"])
    header_font = Font(name=style_cfg["font_name"], size=style_cfg["font_size"], bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=style_cfg["header_color"], end_color=style_cfg["header_color"], fill_type="solid")

    headers = ["File", "Page", "Flavor", "Backend", "Rows", "Columns", "Accuracy", "Whitespace", "Sheet"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for r, (rec, sheet_name) in enumerate(zip(records, sheet_names), start=2):
        values = [
            rec["file"], rec["page"], rec["flavor"], rec["backend"] or "n/a",
            rec["shape"][0], rec["shape"][1], round(rec["accuracy"], 1),
            round(rec["whitespace"], 1), sheet_name,
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v).font = base_font

    last_row = len(records) + 1
    if last_row >= 2:
        ref = f"A1:{get_column_letter(len(headers))}{last_row}"
        table_style_name = TABLE_STYLE_PRESETS.get(style_cfg["style_preset"])
        if table_style_name:
            tbl = Table(displayName="Summary_Tbl", ref=ref)
            tbl.tableStyleInfo = TableStyleInfo(
                name=table_style_name, showRowStripes=True, showColumnStripes=False,
            )
            ws.add_table(tbl)

    if style_cfg["autofit"]:
        autofit_columns(ws, 1, 1, last_row, len(headers))
    ws.freeze_panes = "A2"


def run_pipeline(cfg, log):
    pdf_paths = collect_pdf_paths(cfg["input_path"], cfg["recursive"])
    if not pdf_paths:
        log("No PDF files found for the given input.")
        return None

    all_records = []
    for pdf_path in pdf_paths:
        log(f"Processing {os.path.basename(pdf_path)} ...")
        try:
            records = extract_tables_for_pdf(pdf_path, cfg, log)
            all_records.extend(records)
        except Exception as e:
            log(f"  failed: {e}")

    if not all_records:
        log("No tables extracted from any file.")
        return None

    wb = Workbook()
    wb.remove(wb.active)

    style_cfg = {
        "font_name": cfg["font_name"],
        "font_size": cfg["font_size"],
        "header_color": cfg["header_color"],
        "style_preset": cfg["style_preset"],
        "autofit": cfg["autofit"],
    }

    existing_names = set()
    sheet_names = []
    for rec in all_records:
        base_name = f"{os.path.splitext(rec['file'])[0]}_p{rec['page']}_{rec['flavor']}"
        sheet_name = write_table_sheet(wb, base_name, rec, style_cfg, existing_names)
        sheet_names.append(sheet_name)

    write_summary_sheet(wb, all_records, sheet_names, style_cfg)

    out_path = cfg["output_path"]
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    log(f"Saved {len(all_records)} table(s) to {out_path}")
    return out_path


def default_config():
    return {
        "input_path": "",
        "output_path": "",
        "recursive": False,
        "pages": "all",
        "flavor": "lattice",
        "try_all_flavors": False,
        "backend": "pdfium",
        "use_fallback": True,
        "engine": "combined",
        "filter_pages_enabled": False,
        "filter_text": "",
        "filter_case_sensitive": False,
        "derive_region_enabled": False,
        "header_text": "",
        "footer_text": "",
        "derive_case_sensitive": False,
        "strip_text": "",
        "replace_text": "",
        "split_text": False,
        "flag_size": False,
        "stitch_enabled": True,
        "stitch_match": "column_count",
        "stitch_keep_first_header": False,
        "table_areas": "",
        "table_regions": "",
        "columns": "",
        "row_tol": 2,
        "column_tol": 0,
        "edge_tol": None,
        "line_scale": 15,
        "process_background": False,
        "line_tol": 2,
        "joint_tol": 2,
        "threshold_blocksize": 15,
        "threshold_constant": -2,
        "iterations": 0,
        "erode_iterations": 0,
        "resolution": 300,
        "structure_model": "microsoft/table-transformer-structure-recognition-v1.1-all",
        "detection_model": "microsoft/table-transformer-detection",
        "device": "cpu",
        "detection_threshold": 0.5,
        "structure_threshold": 0.5,
        "crop_padding": 10,
        "ocr": "auto",
        "password": "",
        "parallel": False,
        "cpu_count": None,
        "font_name": "Calibri",
        "font_size": 11,
        "header_color": "4472C4",
        "style_preset": "Blue / White",
        "autofit": True,
    }
